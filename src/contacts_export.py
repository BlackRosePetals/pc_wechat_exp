"""通讯录导出 —— 联系人 / 群聊 → xlsx / csv / html。

命令行（``export -m contacts``）与 Web 通讯录页
（``GET /api/address-book/export``）共用本模块，保证两处导出的列与取值一致。

列定义与取值刻意与旧版 Web CSV 端点保持兼容（epoch 时间戳、Y/N、男/女 标签），
免得下游解析错位；xlsx / html 以可读性优先：消息数存数字（可排序求和）、
时间格式化为 ``YYYY-MM-DD HH:MM:SS``。

数据来源统一为 ``engine.services.address_book.get_all_contacts()``，
筛选统一走 ``engine.services.address_book.filter_contacts()``。
"""
import csv
import io
import os
import re
from datetime import datetime

from engine.constants import TZ
from engine.services import contact_extra

# XML 1.0 不允许的控制字符 —— openpyxl 写单元格时直接抛 IllegalCharacterError。
# 本机 24439 条真实通讯录的签名/描述字段里确实存在这类字节，纯合成数据测不出来。
_ILLEGAL_XML_RE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')

# Excel 单元格上限 32767 字符，超出会让文件在 Excel 里打不开
MAX_CELL_CHARS = 32767

# 支持的输出格式；顺序即 CLI / 文档中的展示顺序
FORMATS = ('xlsx', 'csv', 'html')

MIMETYPES = {
    'csv': 'text/csv; charset=utf-8',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'html': 'text/html; charset=utf-8',
}

# 前 13 列与旧版 /api/address-book/export 的硬编码 CSV 完全一致，勿调整顺序；
# 新列一律追加在末尾（labels 来自 contact.extra_buffer 的标签 id 串）。
COLUMNS = [
    'wxid', 'display_name', 'remark', 'nick_name', 'alias',
    'phone', 'sex', 'region', 'signature', 'description',
    'msg_count', 'last_msg_time', 'is_group',
    'labels',
]

# 性别标签的唯一来源在 contact_extra（extra_buffer field 2），这里转发以便复用
sex_label = contact_extra.sex_label

KIND_LABELS = {
    'all': '联系人与群聊',
    'contacts': '仅联系人',
    'groups': '仅群聊',
}


def _normalize_format(fmt) -> str:
    fmt = str(fmt or 'xlsx').strip().lower().lstrip('.')
    if fmt not in FORMATS:
        raise ValueError('不支持的导出格式: %s（可选: %s）' % (fmt, ' / '.join(FORMATS)))
    return fmt


def region_of(contact: dict) -> str:
    """地区：优先直接给 region，否则拼接 country/province/city 中非空的部分。"""
    direct = (contact.get('region') or '').strip()
    if direct:
        return direct
    return ' '.join(filter(None, [
        (contact.get(k) or '').strip() for k in ('country', 'province', 'city')
    ]))


def clean_text(value) -> str:
    """清洗单个文本值，使其对三种输出格式都安全。

    - 去掉 XML 1.0 非法控制字符（否则 openpyxl 抛 IllegalCharacterError）；
    - 去掉落单代理字符（否则 ``.encode('utf-8')`` 抛 UnicodeEncodeError）；
    - 截断到 Excel 单元格上限（超过 32767 字符 Excel 打不开文件）。

    这些字符在界面上本来就不可见，去掉不影响可读性。
    """
    if not isinstance(value, str):
        return value
    value = _ILLEGAL_XML_RE.sub('', value)
    value = value.encode('utf-8', 'ignore').decode('utf-8')
    if len(value) > MAX_CELL_CHARS:
        value = value[:MAX_CELL_CHARS]
    return value


def build_rows(contacts: list) -> list:
    """联系人 → 表格行（CSV 语义，向后兼容旧端点）。"""
    rows = []
    for c in contacts:
        rows.append([clean_text(v) for v in (
            c.get('wxid', ''),
            c.get('display_name', ''),
            c.get('remark', ''),
            c.get('nick_name', ''),
            c.get('alias', ''),
            c.get('phone', ''),
            sex_label(c.get('sex')),
            region_of(c),
            c.get('signature', ''),
            c.get('description', ''),
            c.get('msg_count') or 0,
            c.get('last_msg_time') or '',
            'Y' if c.get('is_group') else 'N',
            ','.join(c.get('labels') or []),
        )])
    return rows


def _readable_time(ts) -> str:
    if not ts:
        return ''
    try:
        return datetime.fromtimestamp(int(ts), TZ).strftime('%Y-%m-%d %H:%M:%S')
    except (TypeError, ValueError, OSError, OverflowError):
        return ''


def _readable_rows(contacts: list) -> list:
    """xlsx / html 用的取值：消息数为数字、时间为可读字符串。"""
    idx_count = COLUMNS.index('msg_count')
    idx_time = COLUMNS.index('last_msg_time')
    rows = build_rows(contacts)
    for row, c in zip(rows, contacts):
        row[idx_count] = c.get('msg_count') or 0
        row[idx_time] = _readable_time(c.get('last_msg_time'))
    return rows


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------

def render_csv(contacts: list) -> bytes:
    """CSV（UTF-8 带 BOM）。

    BOM 是刻意加的：旧端点的无 BOM 输出在 Excel 里双击打开是乱码。
    """
    buf = io.StringIO(newline='')
    writer = csv.writer(buf)
    writer.writerow(COLUMNS)
    writer.writerows(build_rows(contacts))
    return buf.getvalue().encode('utf-8-sig')


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------

def _display_width(text) -> int:
    """按显示宽度估算列宽：CJK / 全角字符占 2 列。"""
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in str(text))


def _autosize(ws, min_width: int = 8, max_width: int = 60) -> None:
    from openpyxl.utils import get_column_letter
    for col_idx, col_cells in enumerate(ws.iter_cols(), start=1):
        width = min_width
        for cell in col_cells:
            if cell.value is None:
                continue
            width = max(width, _display_width(cell.value) + 2)
            if width >= max_width:      # 已达上限，无需再量剩下的几万行
                break
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width, max_width)


def render_xlsx(contacts: list) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as e:  # pragma: no cover - 依赖已在 requirements 中声明
        raise RuntimeError('导出 xlsx 需要 openpyxl：pip install openpyxl') from e

    wb = Workbook()
    ws = wb.active
    ws.title = '通讯录'
    ws.append(COLUMNS)
    for row in _readable_rows(contacts):
        ws.append(row)

    header_fill = PatternFill('solid', fgColor='1F4E79')
    header_font = Font(bold=True, color='FFFFFF')
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.freeze_panes = 'A2'

    _autosize(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# HTML
# --------------------------------------------------------------------------

def _escape_html(text) -> str:
    return (str(text)
            .replace('&', '&amp;')
            .replace('<', '&lt;')
            .replace('>', '&gt;')
            .replace('"', '&quot;')
            .replace("'", '&#x27;'))


_HTML_STYLE = """
:root { color-scheme: light; }
body { margin: 0; padding: 24px; background: #f6f8fa; color: #24292f;
       font-family: "Microsoft YaHei", "PingFang SC", Segoe UI, sans-serif; }
h1 { font-size: 20px; margin: 0 0 6px 0; }
.meta { color: #57606a; font-size: 13px; margin: 0 0 18px 0; }
.meta b { color: #0969da; }
.table-wrap { overflow-x: auto; background: #fff; border: 1px solid #d0d7de;
              border-radius: 8px; }
table { border-collapse: collapse; width: 100%; font-size: 13px; }
th, td { padding: 7px 10px; text-align: left; white-space: nowrap;
         border-bottom: 1px solid #eaeef2; }
th { background: #1f4e79; color: #fff; font-weight: 600; position: sticky;
     top: 0; }
tbody tr:nth-child(even) { background: #f6f8fa; }
tbody tr:hover { background: #ddf4ff; }
tr.is-group td:nth-child(2) { font-weight: 600; }
.empty { padding: 32px; text-align: center; color: #57606a; }
"""


def render_html(contacts: list, title: str = '微信通讯录', source: str = '') -> str:
    """自包含单文件 HTML（内联样式、无外部资源、离线可看）。"""
    rows = _readable_rows(contacts)
    generated = datetime.now(TZ).strftime('%Y-%m-%d %H:%M:%S')

    head = ''.join('<th>%s</th>' % _escape_html(col) for col in COLUMNS)
    body = []
    for row in rows:
        is_group = str(row[-1]) == 'Y'
        cells = ''.join('<td>%s</td>' % _escape_html(v if v is not None else '')
                        for v in row)
        body.append('<tr class="is-group">%s</tr>' % cells if is_group
                    else '<tr>%s</tr>' % cells)

    if body:
        table = ('<div class="table-wrap"><table><thead><tr>%s</tr></thead>'
                 '<tbody>%s</tbody></table></div>' % (head, ''.join(body)))
    else:
        table = ('<div class="table-wrap"><table><thead><tr>%s</tr></thead>'
                 '</table><div class="empty">没有匹配的通讯录记录</div></div>' % head)

    meta_bits = ['共 <b>%d</b> 条记录' % len(rows),
                 '导出时间 %s' % generated]
    if source:
        meta_bits.append('来源 %s' % _escape_html(source))

    return ('<!DOCTYPE html>\n'
            '<html lang="zh-CN">\n<head>\n'
            '<meta charset="utf-8">\n'
            '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
            '<title>%s</title>\n'
            '<style>%s</style>\n'
            '</head>\n<body>\n'
            '<h1>%s</h1>\n'
            '<p class="meta">%s</p>\n'
            '%s\n'
            '</body>\n</html>\n'
            % (_escape_html(title), _HTML_STYLE, _escape_html(title),
               ' · '.join(meta_bits), table))


# --------------------------------------------------------------------------
# 统一入口
# --------------------------------------------------------------------------

def render(contacts: list, fmt: str = 'xlsx', **kwargs) -> bytes:
    """按格式渲染为字节串（三种格式统一返回 bytes，便于直接写文件/回响应）。"""
    fmt = _normalize_format(fmt)
    if fmt == 'csv':
        return render_csv(contacts)
    if fmt == 'xlsx':
        return render_xlsx(contacts)
    return render_html(contacts, **kwargs).encode('utf-8')


def default_filename(fmt: str = 'xlsx', kind: str = 'all', now=None) -> str:
    fmt = _normalize_format(fmt)
    now = now or datetime.now(TZ)
    stem = 'groups' if kind == 'groups' else 'contacts'
    return '%s_%s.%s' % (stem, now.strftime('%Y%m%d_%H%M%S'), fmt)


def export_contacts(decrypted_dir, out_path: str, fmt: str = 'xlsx', kind: str = 'all',
                    q: str = '', has_chat=None, letter: str = '', sort: str = 'name',
                    label: str = None, contacts: list = None, print_fn=None) -> dict:
    """导出通讯录到 ``out_path``（完整文件路径，父目录自动创建）。

    ``contacts`` 传入时直接用该列表（便于测试与调用方复用已扫描结果），
    否则从 ``decrypted_dir`` 读取。

    Returns: {'path', 'count', 'format', 'bytes'}
    """
    from engine.services.address_book import filter_contacts, get_all_contacts

    if print_fn is None:
        print_fn = print
    fmt = _normalize_format(fmt)
    if kind not in KIND_LABELS:
        raise ValueError('不支持的导出范围: %s（可选: %s）'
                         % (kind, ' / '.join(KIND_LABELS)))

    if contacts is None:
        contacts = get_all_contacts(decrypted_dir)
    contacts = filter_contacts(contacts, q=q, sort=sort, has_chat=has_chat,
                               letter=letter, kind=kind, label=label)

    data = render(contacts, fmt)
    out_dir = os.path.dirname(os.path.abspath(out_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    with open(out_path, 'wb') as f:
        f.write(data)

    print_fn('导出完成: %d 条记录 → %s' % (len(contacts), out_path))
    return {'path': out_path, 'count': len(contacts), 'format': fmt,
            'bytes': len(data)}
